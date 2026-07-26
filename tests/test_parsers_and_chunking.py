from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from app.chunking.parent_child import ParentChildChunker
from app.parsers.registry import ParserRegistry


def _minimal_text_pdf(text: str) -> bytes:
    """Build a dependency-free one-page PDF fixture."""
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = "BT /F1 12 Tf 72 720 Td ({}) Tj ET".format(escaped).encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend("{} 0 obj\n".format(index).encode("ascii"))
        output.extend(obj)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend("xref\n0 {}\n".format(len(objects) + 1).encode("ascii"))
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend("{:010d} 00000 n \n".format(offset).encode("ascii"))
    output.extend(
        (
            "trailer\n<< /Size {} /Root 1 0 R >>\nstartxref\n{}\n%%EOF\n".format(
                len(objects) + 1, xref
            )
        ).encode("ascii")
    )
    return bytes(output)


class ParsingAndChunkingTest(unittest.TestCase):
    def test_markdown_sections_are_injected_into_child_chunks(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "sample.md"
            path.write_text(
                "# Retrieval\n" + ("Dense retrieval content. " * 60) + "\n"
                "## Evaluation\nRecall and MRR are measured.",
                encoding="utf-8",
            )
            elements = ParserRegistry().parse(path)
            chunks = ParentChildChunker(parent_size=500, child_size=180, child_overlap=30).chunk(
                "doc1", elements
            )

        children = [item for item in chunks if item.level == "child"]
        self.assertTrue(children)
        self.assertTrue(any("[Section: Retrieval]" in item.content for item in children))
        self.assertTrue(any(item.section == "Evaluation" for item in children))
        self.assertTrue(all(item.parent_id for item in children))

    def test_unsupported_extension_is_rejected(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "unsafe.exe"
            path.write_bytes(b"MZ")
            with self.assertRaises(ValueError):
                ParserRegistry().parse(path)

    def test_docx_heading_and_table_are_preserved(self) -> None:
        from docx import Document

        with TemporaryDirectory() as directory:
            path = Path(directory) / "policy.docx"
            document = Document()
            document.add_heading("Security", level=1)
            document.add_paragraph("Secrets must use environment variables.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Risk"
            table.cell(0, 1).text = "Action"
            table.cell(1, 0).text = "Missing evidence"
            table.cell(1, 1).text = "Abstain"
            document.save(str(path))
            elements = ParserRegistry().parse(path)

        self.assertTrue(any(item.content_type == "heading" for item in elements))
        self.assertTrue(any(item.content_type == "table" for item in elements))
        self.assertTrue(any(item.section == "Security" for item in elements))

    def test_xlsx_sheet_is_preserved_as_table(self) -> None:
        from openpyxl import Workbook

        with TemporaryDirectory() as directory:
            path = Path(directory) / "metrics.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Retrieval"
            sheet.append(["Method", "Recall@5"])
            sheet.append(["BM25", 0.5])
            workbook.save(str(path))
            elements = ParserRegistry().parse(path)

        self.assertEqual(len(elements), 1)
        self.assertEqual(elements[0].content_type, "table")
        self.assertEqual(elements[0].section, "Retrieval")
        self.assertIn("Recall@5", elements[0].content)

    def test_pdf_text_and_page_number_are_preserved(self) -> None:
        with TemporaryDirectory() as directory:
            path = Path(directory) / "evidence.pdf"
            path.write_bytes(_minimal_text_pdf("Evidence on page one"))
            elements = ParserRegistry().parse(path)

        self.assertEqual(len(elements), 1)
        self.assertEqual(elements[0].page, 1)
        self.assertIn("Evidence on page one", elements[0].content)


if __name__ == "__main__":
    unittest.main()
