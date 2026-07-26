"""DOCX and XLSX parsers."""

from pathlib import Path
from typing import List, Optional

from app.parsers.base import DocumentParseError, DocumentParser
from app.schemas.documents import ParsedElement


class DocxParser(DocumentParser):
    def parse(self, path: Path) -> List[ParsedElement]:
        try:
            from docx import Document
        except ImportError as exc:
            raise DocumentParseError("DOCX support requires python-docx") from exc

        try:
            document = Document(str(path))
            elements: List[ParsedElement] = []
            section: Optional[str] = None
            for paragraph in document.paragraphs:
                text = paragraph.text.strip()
                if not text:
                    continue
                style = (paragraph.style.name if paragraph.style else "").lower()
                if style.startswith("heading"):
                    section = text
                    elements.append(
                        ParsedElement(content=text, content_type="heading", section=section)
                    )
                else:
                    elements.append(ParsedElement(content=text, section=section))

            for table_index, table in enumerate(document.tables, 1):
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                markdown = _rows_to_markdown(rows)
                if markdown:
                    elements.append(
                        ParsedElement(
                            content=markdown,
                            content_type="table",
                            section=section,
                            metadata={"table_index": table_index},
                        )
                    )
            return elements
        except Exception as exc:
            raise DocumentParseError("Failed to parse DOCX: {}".format(exc)) from exc


class XlsxParser(DocumentParser):
    def parse(self, path: Path) -> List[ParsedElement]:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise DocumentParseError("XLSX support requires openpyxl") from exc

        try:
            workbook = load_workbook(str(path), read_only=True, data_only=True)
            elements: List[ParsedElement] = []
            for sheet in workbook.worksheets:
                rows = [
                    ["" if value is None else str(value) for value in row]
                    for row in sheet.iter_rows(values_only=True)
                ]
                markdown = _rows_to_markdown(rows)
                if markdown:
                    elements.append(
                        ParsedElement(
                            content=markdown,
                            content_type="table",
                            section=sheet.title,
                            metadata={"sheet": sheet.title},
                        )
                    )
            workbook.close()
            return elements
        except Exception as exc:
            raise DocumentParseError("Failed to parse XLSX: {}".format(exc)) from exc


def _rows_to_markdown(rows: List[List[str]]) -> str:
    cleaned = [row for row in rows if any(str(value).strip() for value in row)]
    if not cleaned:
        return ""
    width = max(len(row) for row in cleaned)
    normalized = [row + [""] * (width - len(row)) for row in cleaned]
    escaped = [
        [str(value).replace("|", "\\|").replace("\n", " ") for value in row]
        for row in normalized
    ]
    header = escaped[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    lines.extend("| " + " | ".join(row) + " |" for row in escaped[1:])
    return "\n".join(lines)
